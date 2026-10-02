#!/usr/bin/env python3
"""
BARKLY DOCX
DOCX → Barkly-style HTML documentation generator.

Accepts either:
    python barkly_docx.py document.docx
or:
    python barkly_docx.py ./documents

When given a folder, all .docx files are discovered recursively and
converted while preserving their directory structure.

Example:

    documents/
    ├── requirements.docx
    ├── research/
    │   └── paper.docx
    └── architecture/
        └── design.docx

becomes:

    docs/
    ├── requirements/
    │   └── index.html
    ├── research/
    │   └── paper/
    │       └── index.html
    └── architecture/
        └── design/
            └── index.html

Dependency:
    pip install python-docx
"""

from __future__ import annotations

import argparse
import html
import re
from pathlib import Path

from docx import Document

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def escape(value: str) -> str:
    """Safely escape text for HTML."""
    return html.escape(value or "")


def slugify(value: str) -> str:
    """Turn a document title into a safe directory name."""
    value = value.strip().lower()
    value = re.sub(r"[^\w\s-]", "", value)
    value = re.sub(r"[-\s]+", "-", value)
    return value.strip("-_") or "document"


def document_title(document: Document, source: Path) -> str:
    """
    Determine the document title.

    Preference:
        1. First Heading 1
        2. Word document title metadata
        3. Filename
    """

    for paragraph in document.paragraphs:
        if paragraph.style and paragraph.style.name:
            if paragraph.style.name.lower() in {
                "heading 1",
                "title",
            }:
                text = paragraph.text.strip()
                if text:
                    return text

    metadata_title = document.core_properties.title

    if metadata_title and metadata_title.strip():
        return metadata_title.strip()

    return source.stem.replace("_", " ").replace("-", " ").title()


def paragraph_class(style_name: str) -> str:
    """Return a CSS class for a paragraph style."""
    return slugify(style_name or "normal")


# ---------------------------------------------------------------------------
# HTML rendering
# ---------------------------------------------------------------------------


def render_paragraph(paragraph) -> str:
    """Convert one DOCX paragraph into HTML."""

    text = paragraph.text.strip()

    if not text:
        return ""

    style_name = paragraph.style.name if paragraph.style else "Normal"
    style_lower = style_name.lower()

    if style_lower == "title":
        return f'<h1 class="document-title">{escape(text)}</h1>'

    if style_lower.startswith("heading 1"):
        return f"<h1>{escape(text)}</h1>"

    if style_lower.startswith("heading 2"):
        return f"<h2>{escape(text)}</h2>"

    if style_lower.startswith("heading 3"):
        return f"<h3>{escape(text)}</h3>"

    if style_lower.startswith("heading 4"):
        return f"<h4>{escape(text)}</h4>"

    # Lists are handled separately by render_blocks().
    if "list bullet" in style_lower:
        return f"<li>{escape(text)}</li>"

    if "list number" in style_lower:
        return f"<li>{escape(text)}</li>"

    return (
        f'<p class="paragraph {paragraph_class(style_name)}">' f"{escape(text)}" f"</p>"
    )


def render_table(table) -> str:
    """Convert a DOCX table into HTML."""

    rows = table.rows

    if not rows:
        return ""

    output = ['<div class="table-wrap">']
    output.append("<table>")

    for row_index, row in enumerate(rows):
        output.append("<tr>")

        for cell in row.cells:
            cell_text = cell.text.strip()

            if row_index == 0:
                output.append(f"<th>{escape(cell_text)}</th>")
            else:
                output.append(f"<td>{escape(cell_text)}</td>")

        output.append("</tr>")

    output.append("</table>")
    output.append("</div>")

    return "\n".join(output)


def render_blocks(document: Document) -> str:
    """
    Render paragraphs and tables in their original document order.

    Consecutive bullet paragraphs become one <ul>.
    Consecutive numbered paragraphs become one <ol>.
    """

    output: list[str] = []

    paragraphs = iter(document.paragraphs)
    tables = iter(document.tables)

    paragraph_map = {}
    table_map = {}

    for paragraph in document.paragraphs:
        paragraph_map[id(paragraph._p)] = paragraph

    for table in document.tables:
        table_map[id(table._tbl)] = table

    current_list: list[str] = []
    current_list_type: str | None = None

    def flush_list():
        nonlocal current_list, current_list_type

        if not current_list:
            return

        tag = current_list_type or "ul"

        output.append(f"<{tag}>\n" + "\n".join(current_list) + f"\n</{tag}>")

        current_list = []
        current_list_type = None

    # Walk the actual Word body so paragraphs and tables stay ordered.
    body = document.element.body

    for element in body.iterchildren():

        # ---------------------------------------------------------------
        # Paragraph
        # ---------------------------------------------------------------

        if element.tag.endswith("}p"):
            paragraph = paragraph_map.get(id(element))

            if paragraph is None:
                continue

            text = paragraph.text.strip()

            if not text:
                flush_list()
                continue

            style_name = paragraph.style.name if paragraph.style else ""
            style_lower = style_name.lower()

            if "list bullet" in style_lower:
                if current_list_type not in (None, "ul"):
                    flush_list()

                current_list_type = "ul"
                current_list.append(f"<li>{escape(text)}</li>")
                continue

            if "list number" in style_lower:
                if current_list_type not in (None, "ol"):
                    flush_list()

                current_list_type = "ol"
                current_list.append(f"<li>{escape(text)}</li>")
                continue

            flush_list()
            output.append(render_paragraph(paragraph))

        # ---------------------------------------------------------------
        # Table
        # ---------------------------------------------------------------

        elif element.tag.endswith("}tbl"):
            flush_list()

            table = table_map.get(id(element))

            if table is not None:
                output.append(render_table(table))

    flush_list()

    return "\n".join(block for block in output if block.strip())


# ---------------------------------------------------------------------------
# HTML document
# ---------------------------------------------------------------------------


def render_html(
    title: str,
    source: Path,
    content: str,
) -> str:
    """Build the complete Barkly HTML document."""

    safe_title = escape(title)
    safe_source = escape(str(source))

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta
    name="viewport"
    content="width=device-width, initial-scale=1.0"
>

<title>{safe_title} — Barkly Docs</title>

<style>

:root {{
    --bg: #090909;
    --panel: #111111;
    --panel-2: #151515;
    --text: #f2f2f2;
    --muted: #8d8d8d;
    --line: #292929;
    --accent: #ff4fa3;
    --accent-soft: rgba(255, 79, 163, 0.12);
}}

* {{
    box-sizing: border-box;
}}

html {{
    scroll-behavior: smooth;
}}

body {{
    margin: 0;
    background: var(--bg);
    color: var(--text);
    font-family:
        Inter,
        ui-sans-serif,
        system-ui,
        -apple-system,
        BlinkMacSystemFont,
        "Segoe UI",
        sans-serif;
    line-height: 1.7;
}}

a {{
    color: var(--accent);
}}

.topbar {{
    position: sticky;
    top: 0;
    z-index: 20;

    display: flex;
    align-items: center;
    justify-content: space-between;

    padding: 18px 28px;

    background: rgba(9, 9, 9, 0.92);
    border-bottom: 1px solid var(--line);

    backdrop-filter: blur(12px);
}}

.brand {{
    color: var(--text);
    text-decoration: none;
    font-size: 13px;
    font-weight: 800;
    letter-spacing: 0.16em;
}}

.brand-symbol {{
    color: var(--accent);
    margin-right: 8px;
}}

.status {{
    color: var(--muted);
    font-size: 11px;
    letter-spacing: 0.12em;
}}

.status-dot {{
    display: inline-block;
    width: 7px;
    height: 7px;
    margin-right: 7px;

    border-radius: 50%;
    background: var(--accent);
}}

.hero {{
    position: relative;
    overflow: hidden;

    padding: 100px 28px 70px;

    border-bottom: 1px solid var(--line);
}}

.hero::before {{
    content: "";

    position: absolute;
    inset: 0;

    background:
        linear-gradient(
            rgba(255, 79, 163, 0.04) 1px,
            transparent 1px
        ),
        linear-gradient(
            90deg,
            rgba(255, 79, 163, 0.04) 1px,
            transparent 1px
        );

    background-size: 40px 40px;

    mask-image: linear-gradient(
        to bottom,
        black,
        transparent
    );
}}

.hero-inner {{
    position: relative;

    width: min(1100px, 100%);
    margin: auto;
}}

.kicker {{
    color: var(--accent);
    font-size: 11px;
    font-weight: 800;
    letter-spacing: 0.2em;
}}

.hero h1 {{
    max-width: 900px;

    margin: 18px 0;

    font-size: clamp(42px, 8vw, 92px);
    line-height: 0.95;
    letter-spacing: -0.055em;
}}

.hero p {{
    max-width: 700px;

    color: var(--muted);
    font-size: 17px;
}}

.source {{
    margin-top: 28px;

    color: #666;
    font-family: monospace;
    font-size: 11px;
}}

.document {{
    width: min(1000px, calc(100% - 40px));

    margin: 50px auto 100px;
}}

.document-card {{
    padding: clamp(28px, 5vw, 60px);

    background: var(--panel);

    border: 1px solid var(--line);
}}

.document-card h1,
.document-card h2,
.document-card h3,
.document-card h4 {{
    color: var(--text);
    line-height: 1.2;
}}

.document-card h1 {{
    margin-top: 55px;
    font-size: clamp(30px, 5vw, 48px);
}}

.document-card h2 {{
    margin-top: 48px;
    font-size: clamp(24px, 4vw, 36px);
}}

.document-card h3 {{
    margin-top: 36px;
    font-size: 25px;
}}

.document-card h4 {{
    margin-top: 28px;
    font-size: 20px;
}}

.document-card p {{
    max-width: 850px;
    color: #c8c8c8;
}}

.document-card ul,
.document-card ol {{
    max-width: 850px;
    padding-left: 25px;
    color: #c8c8c8;
}}

.document-card li {{
    margin: 8px 0;
}}

.table-wrap {{
    width: 100%;
    overflow-x: auto;

    margin: 30px 0;
}}

table {{
    width: 100%;
    min-width: 500px;

    border-collapse: collapse;

    background: var(--panel-2);
}}

th,
td {{
    padding: 14px 16px;

    text-align: left;
    vertical-align: top;

    border: 1px solid var(--line);
}}

th {{
    color: var(--accent);
    font-size: 12px;
    letter-spacing: 0.08em;
    text-transform: uppercase;
}}

td {{
    color: #c8c8c8;
}}

footer {{
    padding: 30px 28px;

    border-top: 1px solid var(--line);

    color: #666;
    font-family: monospace;
    font-size: 11px;
    text-align: center;
}}

@media (max-width: 700px) {{

    .topbar {{
        padding: 15px 18px;
    }}

    .status {{
        display: none;
    }}

    .hero {{
        padding: 70px 20px 50px;
    }}

    .document {{
        width: calc(100% - 24px);
        margin-top: 25px;
    }}

    .document-card {{
        padding: 24px 20px;
    }}

}}

</style>
</head>

<body>

<header class="topbar">

    <a class="brand" href="/">
        <span class="brand-symbol">◆</span>
        BARKLY LABS
    </a>

    <div class="status">
        <span class="status-dot"></span>
        DOCUMENTATION
    </div>

</header>


<section class="hero">

    <div class="hero-inner">

        <div class="kicker">
            BARKLY / KNOWLEDGE
        </div>

        <h1>
            {safe_title}
        </h1>

        <p>
            Generated documentation produced by Barkly Docs.
        </p>

        <div class="source">
            SOURCE: {safe_source}
        </div>

    </div>

</section>


<main class="document">

    <article class="document-card">

        {content}

    </article>

</main>


<footer>
    BARKLY DOCS · DOCUMENTATION IS INFRASTRUCTURE.
</footer>

</body>
</html>
"""


# ---------------------------------------------------------------------------
# Conversion
# ---------------------------------------------------------------------------


def convert_file(
    source: Path,
    output_root: Path,
    input_root: Path | None = None,
) -> Path:
    """Convert one DOCX file."""

    document = Document(source)

    title = document_title(document, source)
    content = render_blocks(document)

    # ---------------------------------------------------------------
    # Determine output directory.
    #
    # If converting a folder, preserve its relative structure.
    # ---------------------------------------------------------------

    if input_root is not None:
        relative = source.relative_to(input_root)
        output_dir = output_root / relative.parent / slugify(source.stem)
    else:
        output_dir = output_root / slugify(source.stem)

    output_dir.mkdir(parents=True, exist_ok=True)

    output_file = output_dir / "index.html"

    output_file.write_text(
        render_html(
            title=title,
            source=source,
            content=content,
        ),
        encoding="utf-8",
    )

    return output_file


def find_docx_files(root: Path) -> list[Path]:
    """Recursively discover DOCX files."""

    return sorted(
        path
        for path in root.rglob("*.docx")
        if path.is_file() and not path.name.startswith("~$")
    )


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def main() -> None:
    parser = argparse.ArgumentParser(
        description=("Convert DOCX documents into Barkly-style HTML " "documentation.")
    )

    parser.add_argument(
        "input",
        type=Path,
        help="DOCX file or folder containing DOCX files",
    )

    parser.add_argument(
        "--output",
        "-o",
        type=Path,
        default=Path("docs"),
        help="Output directory (default: ./docs)",
    )

    args = parser.parse_args()

    input_path = args.input.resolve()
    output_root = args.output.resolve()

    if not input_path.exists():
        raise SystemExit(f"ERROR: Input does not exist: {input_path}")

    print()
    print("BARKLY DOCS / DOCX CONVERTER")
    print("=" * 45)
    print()

    # ------------------------------------------------------------------
    # Single file
    # ------------------------------------------------------------------

    if input_path.is_file():

        if input_path.suffix.lower() != ".docx":
            raise SystemExit("ERROR: Input file must be a .docx document.")

        print(f"INPUT : {input_path}")
        print(f"OUTPUT: {output_root}")
        print()

        output_file = convert_file(
            input_path,
            output_root,
        )

        print(f"✓ {input_path.name}")
        print(f"  → {output_file}")
        print()
        print("DONE.")
        return

    # ------------------------------------------------------------------
    # Folder
    # ------------------------------------------------------------------

    if input_path.is_dir():

        files = find_docx_files(input_path)

        print(f"INPUT : {input_path}")
        print(f"OUTPUT: {output_root}")
        print()
        print(f"FOUND {len(files)} DOCX DOCUMENT(S)")
        print()

        if not files:
            print("No .docx files found.")
            return

        converted = 0
        failed = 0

        for index, source in enumerate(files, start=1):

            print(f"[{index}/{len(files)}] " f"{source.relative_to(input_path)}")

            try:

                output_file = convert_file(
                    source,
                    output_root,
                    input_root=input_path,
                )

                print(f"        → {output_file}")

                converted += 1

            except Exception as exc:

                print(f"        ERROR: {exc}")

                failed += 1

        print()
        print("=" * 45)
        print("BARKLY DOCS COMPLETE")
        print()
        print(f"Converted : {converted}")
        print(f"Failed    : {failed}")
        print(f"Output    : {output_root}")
        print()

        return

    raise SystemExit("ERROR: Input must be a DOCX file or directory.")


if __name__ == "__main__":
    main()
